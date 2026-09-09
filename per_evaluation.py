"""
PER Evaluation — EN-US / EN-PK / Pipeline Variants vs. Linguist Reference
=========================================================================

Evaluation design
-----------------
There are FOUR hypotheses and ONE gold/reference transcription:

    HYPOTHESES
      1. Raw eSpeak EN-US
      2. Raw eSpeak EN-PK
      3. Proposed pipeline + EN-US fallback
      4. Proposed pipeline + EN-PK fallback

    GOLD / REFERENCE
      5. Linguist transcription (Maham)

Maham's transcription is the ONLY reference. All four systems are scored
against it using phoneme error rate (PER).

The proposed pipeline remains lexicon-first and uses the same preprocessing,
routing, and Pakistani-English phonemic collapses. The only difference
between the two pipeline variants is the eSpeak voice used whenever the
lexicon/routing does not provide a pronunciation.

Spreadsheet layout (Evaluation Sentences)
------------------------------------------
A  No.
B  ID
C  Sentence
D  Coverage / Key Words
E  espeak en-us Output
F  espeak en-pk Output
G  Pipeline Output with en-us fallback
H  Pipeline Output with en-pk fallback
I  Reference IPA (Maham)

IMPORTANT:
  --fill regenerates ONLY the four hypothesis columns (E-H).
  Column I is never overwritten, because it is manually supplied by Maham.

Usage
-----
    python per_evaluation_sentences.py --excel Eval.xlsx --lexicon lex.txt --fill
    python per_evaluation_sentences.py --excel Eval.xlsx --lexicon lex.txt
"""

import argparse
import csv
import re
import subprocess
import sys
from collections import OrderedDict

import editdistance
import openpyxl
from num2words import num2words

ESPEAK_PATH = r"C:\Program Files\eSpeak NG\espeak-ng.exe"

# Column indices (1-based) in the Evaluation Sentences sheet
COL_NO = 1
COL_ID = 2
COL_SENT = 3
COL_COVERAGE = 4
COL_ENUS = 5
COL_ENPK = 6
COL_PIPE_US = 7
COL_PIPE_PK = 8
COL_REF = 9


# --------------------------------------------------------------------------- #
# Phoneme inventory — longest-first for greedy tokenisation
# --------------------------------------------------------------------------- #

PHONEME_INVENTORY = sorted([
    # Pakistani English additional phones
    "t\u032A\u02B0", "d\u032A\u02B0", "\u0256\u02B0", "\u027D\u02B0",
    "\u0288\u02B0",  # ʈʰ - aspirated voiceless retroflex stop (new: t -> ʈ rule)
    "t\u032A", "d\u032A",
    "p\u02B0", "b\u02B0", "k\u02B0", "\u0261\u02B0",
    "t\u0283\u02B0", "d\u0292\u02B0",
    "m\u02B0", "n\u02B0", "l\u02B0", "r\u02B0", "j\u02B0",
    # Nasalised vowels
    "\u0251\u0303\u02D0", "e\u0303\u02D0", "i\u0303\u02D0",
    "o\u0303\u02D0", "u\u0303\u02D0", "\u0254\u0303\u02D0", "\u00E6\u0303\u02D0",
    "\u026A\u0303", "\u028A\u0303", "\u0259\u0303",
    # Long monophthongs (including collapse targets)
    "\u0251\u02D0", "i\u02D0", "u\u02D0", "\u025C\u02D0",
    "e\u02D0", "o\u02D0", "\u0254\u02D0",
    # Remaining diphthongs
    "a\u026A", "\u0251\u026A", "a\u028A", "\u0251\u028A", "\u0254\u026A",
    # Affricates / postalveolars
    "t\u0283", "d\u0292", "\u0283", "\u0292", "\u014B",
    # Pakistani English consonants
    "\u027D", "\u0256", "\u0288", "q", "\u0294", "x", "\u0263",  # ʈ added (new: t -> ʈ rule)
    # Single-character phones (theta, eth, w excluded — collapsed)
    "p", "b", "k", "\u0261", "m", "n",  # t, d removed - collapsed to ʈ, ɖ
    "f", "v", "s", "z", "h", "\u0279", "j", "l", "\u027E",
    # Vowels
    "\u026A", "\u028A", "\u0259", "\u028C", "\u025B", "\u00E6", "\u0254", "\u0252",
    "\u0250", "\u0268", "\u1D7B", "\u025A", "\u025D",
    # Suprasegmentals & punctuation
    "\u02C8", "\u02CC", "\u02D0",
    ";", ":", ",", ".", "!", "?", "\u2014", "\u2026",
], key=len, reverse=True)


# --------------------------------------------------------------------------- #
# IPA tokenisation & PER
# --------------------------------------------------------------------------- #

def tokenize_ipa(ipa_string, inventory=PHONEME_INVENTORY):
    """Greedy longest-match IPA tokenisation. Unknown chars are kept as-is."""
    tokens, i = [], 0
    s = (ipa_string or "").replace(" ", "")
    while i < len(s):
        for phone in inventory:
            if s.startswith(phone, i):
                tokens.append(phone)
                i += len(phone)
                break
        else:
            tokens.append(s[i])
            i += 1
    return tokens


def phoneme_error_rate(reference_ipa, hypothesis_ipa):
    """Return (PER, n_errors, n_reference_phones)."""
    ref = tokenize_ipa(reference_ipa)
    hyp = tokenize_ipa(hypothesis_ipa)

    if not ref:
        return None, 0, 0

    dist = editdistance.eval(ref, hyp)
    return dist / len(ref), dist, len(ref)


# --------------------------------------------------------------------------- #
# eSpeak
# --------------------------------------------------------------------------- #

def g2p(text, voice="en-us"):
    """Phonemise text with eSpeak NG. Returns empty string on failure."""
    try:
        result = subprocess.run(
            [ESPEAK_PATH, "-q", "--ipa=3", "-v", voice],
            input=text,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            encoding="utf-8",
        )
        if result.returncode != 0:
            print(
                f"  [WARN] espeak {voice} failed on {text!r}: "
                f"{result.stderr.strip()}",
                file=sys.stderr,
            )
            return ""
        return " ".join(result.stdout.split())
    except Exception as exc:
        print(
            f"  [WARN] espeak {voice} failed on {text!r}: {exc}",
            file=sys.stderr,
        )
        return ""


# --------------------------------------------------------------------------- #
# Finalised Pakistani-English phonemic collapses
# --------------------------------------------------------------------------- #

def apply_pake_collapses(ipa):
    """Apply collapses ONLY to eSpeak output, never to lexicon entries.

    Seven rules total (per CLE Urdu Phonetic Inventory + Maham's confirmation):
        theta -> dental t-aspirated   (θ -> t̪ʰ)
        eth   -> dental d             (ð -> d̪)
        w     -> v                    (labiodental)
        eI    -> e:                   (diphthong monophthongisation)
        oU    -> o:                   (diphthong monophthongisation)
        t     -> ʈ                    (Pakistani English has no plain alveolar /t/)
        d     -> ɖ                    (Pakistani English has no plain alveolar /d/)

    The t->ʈ and d->ɖ rules use a negative lookahead so they do NOT fire on
    t̪/t̪ʰ/tʃ/tʃʰ or d̪/d̪ʰ/dʒ/dʒʰ (dental, aspirated-dental, and affricate
    forms are left untouched, including both the tie-bar t͡ʃ/d͡ʒ notation
    and eSpeak's zero-width-joiner t‍ʃ/d‍ʒ notation, since ZWJ is stripped
    on the first line before these rules run).
    """
    ipa = ipa.replace("\u200d", "")  # zero-width joiner (must run first)
    ipa = ipa.replace("e\u026A", "e\u02D0").replace("\u0259\u026A", "e\u02D0")
    ipa = ipa.replace("o\u028A", "o\u02D0").replace("\u0259\u028A", "o\u02D0")
    ipa = ipa.replace("\u03B8", "t\u032A\u02B0")
    ipa = ipa.replace("\u00F0", "d\u032A")
    ipa = ipa.replace("w", "v")
    ipa = re.sub(r"t(?![\u032A\u0283\u0361])", "\u0288", ipa)  # t -> ʈ
    ipa = re.sub(r"d(?![\u032A\u0292\u0361])", "\u0256", ipa)  # d -> ɖ
    return ipa


# --------------------------------------------------------------------------- #
# Routing
# --------------------------------------------------------------------------- #

ABBREVIATIONS = {
    "mrs": "misess", "mr": "mister", "dr": "doctor", "co": "company",
    "jr": "junior", "maj": "major", "gen": "general", "drs": "doctors",
    "rev": "reverend", "lt": "lieutenant", "hon": "honorable",
    "sgt": "sergeant", "capt": "captain", "esq": "esquire",
    "ltd": "limited", "col": "colonel", "ft": "fort", "rs": "rupees",
    "st": "street",
}

ROMAN_NUMERALS = {
    "i": "one", "ii": "two", "iii": "three", "iv": "four", "v": "five",
    "vi": "six", "vii": "seven", "viii": "eight", "ix": "nine", "x": "ten",
}

PUNCTUATION = set(";:,.!?\u00A1\u00BF\u2014\u2026\"\u00AB\u00BB\u201C\u201D")

_TOKEN_RE = re.compile(r"\([ivx]+\)|\w+|[^\w\s]", re.IGNORECASE)
_PAREN_ROMAN_RE = re.compile(
    r"^\((i|ii|iii|iv|v|vi|vii|viii|ix|x)\)$", re.IGNORECASE
)
_NUM_ALPHA_RE = re.compile(r"^(\d+)([a-zA-Z]+)$")
_INTEGER_RE = re.compile(r"^\d+$")
_WORD_RE = re.compile(r"\w+")
_ABBREV_RE = {
    k: re.compile(r"\b%s\." % re.escape(k), re.IGNORECASE)
    for k in ABBREVIATIONS
}


def expand_abbreviations(text):
    for key, pattern in _ABBREV_RE.items():
        text = pattern.sub(ABBREVIATIONS[key], text)
    return text


def analyze_token(token, lexicon, fallback_voice="en-us"):
    """
    Route one token:

        lexicon -> roman -> alphanumeric -> integer -> eSpeak fallback

    Every eSpeak fallback uses the requested pipeline voice. This means
    Pipeline + EN-US and Pipeline + EN-PK differ only in their fallback voice;
    lexicon entries remain identical between the two variants.
    """
    lower = token.lower()

    if lower in lexicon:
        return lexicon[lower]

    roman = _PAREN_ROMAN_RE.match(token)
    if roman:
        return apply_pake_collapses(
            g2p(ROMAN_NUMERALS[roman.group(1).lower()], fallback_voice)
        )

    num_alpha = _NUM_ALPHA_RE.match(lower)
    if num_alpha:
        num_part, alpha_part = num_alpha.groups()

        num_ipa = apply_pake_collapses(
            g2p(num2words(int(num_part)), fallback_voice)
        )

        alpha_ipa = (
            lexicon[alpha_part]
            if alpha_part in lexicon
            else apply_pake_collapses(g2p(alpha_part, fallback_voice))
        )

        return (num_ipa + " " + alpha_ipa).strip()

    if _INTEGER_RE.match(lower):
        return apply_pake_collapses(
            g2p(num2words(int(lower)), fallback_voice)
        )

    if _WORD_RE.match(lower):
        return apply_pake_collapses(g2p(lower, fallback_voice))

    if token in PUNCTUATION:
        return token

    return ""


def pipeline_sentence(sentence, lexicon, fallback_voice="en-us"):
    """
    Run the proposed pipeline.

    The pipeline is lexicon-first. OOV words, numeric expansions, and Roman
    numeral expansions use the selected eSpeak fallback voice.
    """
    normalised = expand_abbreviations(sentence)
    tokens = _TOKEN_RE.findall(normalised)

    parts = [
        analyze_token(tok, lexicon, fallback_voice)
        for tok in tokens
    ]

    return " ".join(p for p in parts if p).strip()


# --------------------------------------------------------------------------- #
# Workbook I/O
# --------------------------------------------------------------------------- #

def load_lexicon(path):
    lexicon = {}
    with open(path, "r", encoding="utf-8") as fh:
        for line in fh:
            if "\t" not in line:
                continue
            word, ipa = line.rstrip("\n").split("\t", 1)
            lexicon[word.strip().lower()] = ipa.strip()
    return lexicon


def read_sheet(excel_path, sheet_name="Evaluation Sentences"):
    """
    Return (workbook, worksheet, rows).

    Maham's reference is read from column I and is NEVER generated by this
    script. Section header rows (no ID) set the current section label.
    """
    wb = openpyxl.load_workbook(excel_path)

    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet {sheet_name!r} not found. Available: {wb.sheetnames}"
        )

    ws = wb[sheet_name]
    rows = []
    section = "Unsectioned"

    for idx, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        cell_a = (
            str(row[COL_NO - 1]).strip()
            if len(row) >= COL_NO and row[COL_NO - 1]
            else ""
        )
        sent_id = (
            str(row[COL_ID - 1]).strip()
            if len(row) >= COL_ID and row[COL_ID - 1]
            else ""
        )
        sentence = (
            str(row[COL_SENT - 1]).strip()
            if len(row) >= COL_SENT and row[COL_SENT - 1]
            else ""
        )

        if cell_a.lower().startswith("section"):
            section = cell_a
            continue

        if not sent_id or not sentence:
            continue

        ref = ""
        if len(row) >= COL_REF and row[COL_REF - 1]:
            ref = str(row[COL_REF - 1]).strip()

        rows.append({
            "row": idx,
            "section": section,
            "id": sent_id,
            "sentence": sentence,
            "ref_ipa": ref,
        })

    return wb, ws, rows


def fill_columns(excel_path, lexicon_path, sheet_name="Evaluation Sentences"):
    """
    Regenerate the FOUR hypothesis columns only.

    Column I (Reference IPA (Maham)) is intentionally preserved.
    """
    lexicon = load_lexicon(lexicon_path)
    wb, ws, rows = read_sheet(excel_path, sheet_name)

    # Keep the spreadsheet structure exactly as designed.
    ws.cell(row=1, column=COL_ENUS, value="espeak en-us Output")
    ws.cell(row=1, column=COL_ENPK, value="espeak en-pk Output")
    ws.cell(row=1, column=COL_PIPE_US,
            value="Pipeline Output with en-us fallback")
    ws.cell(row=1, column=COL_PIPE_PK,
            value="Pipeline Output with en-pk fallback")
    ws.cell(row=1, column=COL_REF, value="Reference IPA (Maham)")

    for rec in rows:
        # FOUR hypotheses
        enus = g2p(rec["sentence"], "en-us")
        enpk = g2p(rec["sentence"], "en-pk")
        pipe_us = pipeline_sentence(rec["sentence"], lexicon, "en-us")
        pipe_pk = pipeline_sentence(rec["sentence"], lexicon, "en-pk")

        ws.cell(row=rec["row"], column=COL_ENUS, value=enus)
        ws.cell(row=rec["row"], column=COL_ENPK, value=enpk)
        ws.cell(row=rec["row"], column=COL_PIPE_US, value=pipe_us)
        ws.cell(row=rec["row"], column=COL_PIPE_PK, value=pipe_pk)

        # DO NOT write to COL_REF.
        print(f"  {rec['id']}  done")

    wb.save(excel_path)
    print(
        f"\nFilled 4 hypothesis columns for {len(rows)} sentences "
        f"-> {excel_path}"
    )
    print("Reference IPA (Maham) column was preserved.")


# --------------------------------------------------------------------------- #
# Evaluation
# --------------------------------------------------------------------------- #

SYSTEMS = OrderedDict([
    ("enus", "EN-US"),
    ("enpk", "EN-PK"),
    ("pipe_us", "Pipeline + EN-US fallback"),
    ("pipe_pk", "Pipeline + EN-PK fallback"),
])


def evaluate(excel_path, lexicon_path, output_path,
             sheet_name="Evaluation Sentences"):
    """
    Score all FOUR hypotheses against Maham's transcription.

    No fallback reference is generated. A missing Maham reference is treated
    as missing gold data and that sentence is skipped with a warning.
    """
    lexicon = load_lexicon(lexicon_path)
    _, _, rows = read_sheet(excel_path, sheet_name)

    if not rows:
        print("[ERROR] No sentences found in sheet.")
        return

    print(
        f"Scoring 4 hypotheses on {len(rows)} sentences "
        f"against Maham's linguist reference...\n"
    )

    results = []

    # Aggregate edit distance / reference-phone counts.
    totals = {
        key: {"err": 0, "n": 0}
        for key in SYSTEMS
    }

    # Number of sentence-level wins/ties is useful for later analysis.
    sentence_wins = {
        key: 0
        for key in SYSTEMS
    }
    sentence_ties = 0

    by_section = OrderedDict()

    skipped = 0

    for rec in rows:
        ref = rec["ref_ipa"]

        # Maham is the ONLY gold reference.
        if not ref:
            skipped += 1
            print(
                f"  [WARN] {rec['id']}: missing Maham reference; "
                f"skipping evaluation for this sentence.",
                file=sys.stderr,
            )
            continue

        # FOUR hypotheses
        enus = g2p(rec["sentence"], "en-us")
        enpk = g2p(rec["sentence"], "en-pk")
        pipe_us = pipeline_sentence(rec["sentence"], lexicon, "en-us")
        pipe_pk = pipeline_sentence(rec["sentence"], lexicon, "en-pk")

        hypotheses = {
            "enus": enus,
            "enpk": enpk,
            "pipe_us": pipe_us,
            "pipe_pk": pipe_pk,
        }

        row_scores = {}

        for key, hypothesis in hypotheses.items():
            per, errors, n_ref = phoneme_error_rate(ref, hypothesis)

            row_scores[key] = {
                "per": per,
                "errors": errors,
                "n_ref": n_ref,
            }

            totals[key]["err"] += errors
            totals[key]["n"] += n_ref

        # Determine sentence-level winner(s): lowest PER.
        valid_scores = {
            key: score["per"]
            for key, score in row_scores.items()
            if score["per"] is not None
        }

        if valid_scores:
            best_per = min(valid_scores.values())
            winners = [
                key for key, value in valid_scores.items()
                if value == best_per
            ]

            if len(winners) == 1:
                sentence_wins[winners[0]] += 1
            else:
                sentence_ties += 1

        sec = rec["section"]
        b = by_section.setdefault(
            sec,
            {
                "count": 0,
                **{
                    key: [0, 0]
                    for key in SYSTEMS
                },
            },
        )

        b["count"] += 1

        for key in SYSTEMS:
            b[key][0] += row_scores[key]["errors"]
            b[key][1] += row_scores[key]["n_ref"]

        results.append({
            "id": rec["id"],
            "section": sec,
            "sentence": rec["sentence"],
            "reference_ipa": ref,

            "enus_ipa": enus,
            "enus_per": round(row_scores["enus"]["per"] * 100, 2),

            "enpk_ipa": enpk,
            "enpk_per": round(row_scores["enpk"]["per"] * 100, 2),

            "pipeline_enus_ipa": pipe_us,
            "pipeline_enus_per": round(
                row_scores["pipe_us"]["per"] * 100, 2
            ),

            "pipeline_enpk_ipa": pipe_pk,
            "pipeline_enpk_per": round(
                row_scores["pipe_pk"]["per"] * 100, 2
            ),

            # Sentence-level winner is useful for "where each system wins".
            "sentence_winner": (
                SYSTEMS[winners[0]]
                if len(winners) == 1
                else "Tie"
            ),
        })

    def pct(err, n):
        return round(err / n * 100, 1) if n else None

    overall = {
        key: pct(totals[key]["err"], totals[key]["n"])
        for key in SYSTEMS
    }

    # ------------------------------------------------------------------ #
    # Console summary
    # ------------------------------------------------------------------ #

    print("\nPER by section (all against Maham):")
    header = f"{'Section':<42}{'N':>5}"
    for key in SYSTEMS:
        header += f"{SYSTEMS[key]:>22}"
    print(header)
    print("-" * (47 + 22 * len(SYSTEMS)))

    for sec, b in by_section.items():
        label = sec if len(sec) <= 40 else sec[:39] + "\u2026"
        line = f"{label:<42}{b['count']:>5}"
        for key in SYSTEMS:
            value = pct(*b[key])
            line += f"{value:>22.1f}" if value is not None else f"{'N/A':>22}"
        print(line)

    print("-" * (47 + 22 * len(SYSTEMS)))

    line = f"{'OVERALL':<42}{len(results):>5}"
    for key in SYSTEMS:
        value = overall[key]
        line += f"{value:>22.1f}" if value is not None else f"{'N/A':>22}"
    print(line)

    print("\nSentence-level wins (lowest PER against Maham):")
    for key in SYSTEMS:
        print(f"  {SYSTEMS[key]:<30}: {sentence_wins[key]}")
    print(f"  {'Ties':<30}: {sentence_ties}")

    if skipped:
        print(
            f"\n[WARN] {skipped} sentence(s) skipped because "
            f"Maham's reference was blank."
        )

    # ------------------------------------------------------------------ #
    # CSV output
    # ------------------------------------------------------------------ #

    fields = [
        "id",
        "section",
        "sentence",
        "reference_ipa",

        "enus_ipa",
        "enus_per",

        "enpk_ipa",
        "enpk_per",

        "pipeline_enus_ipa",
        "pipeline_enus_per",

        "pipeline_enpk_ipa",
        "pipeline_enpk_per",

        "sentence_winner",
    ]

    with open(
        output_path,
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows(results)

        # Overall PER row
        overall_row = {
            "id": "OVERALL",
            "sentence": "Aggregate PER against Maham reference",
        }

        for key in SYSTEMS:
            overall_row[f"{key}_per"] = overall[key]

        writer.writerow(overall_row)

        # Sentence-win counts
        wins_row = {
            "id": "SENTENCE_WINS",
            "sentence": "Number of sentences with lowest PER",
        }

        for key in SYSTEMS:
            wins_row[f"{key}_per"] = sentence_wins[key]

        writer.writerow(wins_row)

    print(f"\nPer-sentence results -> {output_path}")
    print("\nSummary:")
    for key in SYSTEMS:
        print(f"  {SYSTEMS[key]:<30} {overall[key]:.1f}%")
    print(f"\nGold/reference: Maham (linguist transcription)")
    print("All four systems are hypotheses; none is used as the reference.")


# --------------------------------------------------------------------------- #

if __name__ == "__main__":
    ap = argparse.ArgumentParser(
        description=(
            "PER for EN-US, EN-PK, Pipeline+EN-US and Pipeline+EN-PK "
            "against Maham's linguist reference."
        )
    )
    ap.add_argument(
        "--excel",
        required=True,
        help="Evaluation workbook (.xlsx)",
    )
    ap.add_argument(
        "--lexicon",
        required=True,
        help="word<TAB>IPA lexicon",
    )
    ap.add_argument(
        "--output",
        default="per_results_sentences.csv",
        help="CSV output path for evaluation results",
    )
    ap.add_argument(
        "--sheet",
        default="Evaluation Sentences",
        help="Worksheet containing the evaluation sentences",
    )
    ap.add_argument(
        "--fill",
        action="store_true",
        help=(
            "Regenerate the FOUR hypothesis columns (E-H). "
            "Never overwrite Maham's reference in column I."
        ),
    )

    args = ap.parse_args()

    if args.fill:
        fill_columns(args.excel, args.lexicon, args.sheet)
    else:
        evaluate(args.excel, args.lexicon, args.output, args.sheet)
