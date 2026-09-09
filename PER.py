"""
PER Evaluation — Read All 5 Columns Directly (No Recomputation)
==================================================================

Reads the Excel sheet as-is: espeak en-us, espeak en-pk, Pipeline+en-us,
Pipeline+en-pk, and Reference (Maham). Scores all four hypothesis columns
against Reference using PER. Does NOT call espeak, does NOT touch the
lexicon, does NOT recompute anything — whatever is in the cells is what
gets scored, including any manual corrections you've made directly in
the sheet.

Usage
-----
    python per_from_sheet.py --excel "PakE_Evaluation_Sentences.xlsx"
"""

import argparse
import csv
import re
from collections import OrderedDict

import editdistance
import openpyxl

# Column indices (1-based)
COL_NO, COL_ID, COL_SENT, COL_COVERAGE = 1, 2, 3, 4
COL_ENUS, COL_ENPK, COL_PIPE_US, COL_PIPE_PK, COL_REF = 5, 6, 7, 8, 9


# --------------------------------------------------------------------------- #
# Phoneme inventory — longest-first for greedy tokenisation
# --------------------------------------------------------------------------- #

PHONEME_INVENTORY = sorted([
    "t\u032A\u02B0", "d\u032A\u02B0", "\u0288\u02B0",
    "t\u032A", "d\u032A",
    "p\u02B0", "b\u02B0", "k\u02B0", "\u0261\u02B0",
    "t\u0283\u02B0", "d\u0292\u02B0",
    "m\u02B0", "n\u02B0", "l\u02B0", "r\u02B0", "j\u02B0",
    "\u0251\u0303\u02D0", "e\u0303\u02D0", "i\u0303\u02D0",
    "o\u0303\u02D0", "u\u0303\u02D0", "\u0254\u0303\u02D0", "\u00E6\u0303\u02D0",
    "\u026A\u0303", "\u028A\u0303", "\u0259\u0303",
    "\u0251\u02D0", "i\u02D0", "u\u02D0", "\u025C\u02D0",
    "e\u02D0", "o\u02D0", "\u0254\u02D0",
    "a\u026A", "\u0251\u026A", "a\u028A", "\u0251\u028A", "\u0254\u026A",
    "t\u0283", "d\u0292", "\u0283", "\u0292", "\u014B",
    "\u027D", "\u0256", "\u0288", "q", "\u0294", "x", "\u0263",
    "p", "b", "k", "\u0261", "m", "n",
    "f", "v", "s", "z", "h", "\u0279", "j", "l", "\u027E",
    "\u026A", "\u028A", "\u0259", "\u028C", "\u025B", "\u00E6", "\u0254", "\u0252",
    "\u0250", "\u0268", "\u1D7B", "\u025A", "\u025D",
    "\u02C8", "\u02CC", "\u02D0",
    ";", ":", ",", ".", "!", "?", "\u2014", "\u2026", "\u0256\u02B0", "\u027D\u02B0"
], key=len, reverse=True)


def tokenize_ipa(ipa_string, inventory=PHONEME_INVENTORY):
    """Greedy longest-match IPA tokenisation. Unknown chars kept as-is."""
    tokens, i = [], 0
    s = ipa_string.replace(" ", "").replace("\u200d", "")
    while i < len(s):
        for phone in inventory:
            if s.startswith(phone, i):
                tokens.append(phone)
                i += len(phone)
                break
        else:
            tokens.append(s[i])
            i += 1
    return tokens


def phoneme_error_rate(reference_ipa, hypothesis_ipa):
    """Return (per, n_errors, n_reference_phones)."""
    ref = tokenize_ipa(reference_ipa)
    hyp = tokenize_ipa(hypothesis_ipa)
    if not ref:
        return 0.0, 0, 0
    dist = editdistance.eval(ref, hyp)
    return dist / len(ref), dist, len(ref)


# --------------------------------------------------------------------------- #
# Read the sheet
# --------------------------------------------------------------------------- #

def read_sheet(excel_path, sheet_name=None):
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]

    rows, section = [], "Unsectioned"
    for row in ws.iter_rows(min_row=2, values_only=True):
        cell_a = str(row[COL_NO - 1]).strip() if row[COL_NO - 1] else ""
        sent_id = str(row[COL_ID - 1]).strip() if len(row) > 1 and row[COL_ID - 1] else ""
        sentence = str(row[COL_SENT - 1]).strip() if len(row) > 2 and row[COL_SENT - 1] else ""

        if cell_a.lower().startswith("section"):
            section = cell_a
            continue
        if not sent_id or not sentence:
            continue

        def get(col):
            return str(row[col - 1]).strip() if len(row) >= col and row[col - 1] else ""

        rows.append({
            "id": sent_id, "section": section, "sentence": sentence,
            "enus": get(COL_ENUS), "enpk": get(COL_ENPK),
            "pipe_us": get(COL_PIPE_US), "pipe_pk": get(COL_PIPE_PK),
            "ref": get(COL_REF),
        })
    return rows


# --------------------------------------------------------------------------- #
# Evaluate
# --------------------------------------------------------------------------- #

def evaluate(excel_path, output_path, sheet_name=None):
    rows = read_sheet(excel_path, sheet_name)

    scored = [r for r in rows if r["ref"]]
    skipped = [r for r in rows if not r["ref"]]
    if skipped:
        print(f"[INFO] {len(skipped)} rows have no Reference IPA, skipped:")
        for r in skipped:
            print(f"       {r['id']}")

    if not scored:
        print("[ERROR] No rows with Reference IPA found.")
        return

    systems = ["enus", "enpk", "pipe_us", "pipe_pk"]
    totals = {s: {"err": 0, "n": 0} for s in systems}
    by_section = OrderedDict()
    results = []

    for r in scored:
        ref = r["ref"]
        row_result = {"id": r["id"], "section": r["section"], "sentence": r["sentence"],
                      "reference_ipa": ref}

        sec = r["section"]
        bucket = by_section.setdefault(sec, {"count": 0, **{s: [0, 0] for s in systems}})
        bucket["count"] += 1

        for s in systems:
            hyp = r[s]
            per, err, n = phoneme_error_rate(ref, hyp)
            totals[s]["err"] += err
            totals[s]["n"] += n
            bucket[s][0] += err
            bucket[s][1] += n
            row_result[f"{s}_ipa"] = hyp
            row_result[f"{s}_per"] = round(per * 100, 2)

        results.append(row_result)

    def pct(err, n):
        return round(err / n * 100, 1) if n else 0.0

    labels = {"enus": "EN-US", "enpk": "EN-PK", "pipe_us": "PIPE-US", "pipe_pk": "PIPE-PK"}

    print(f"\n{'Section':<42}{'N':>4}" + "".join(f"{labels[s]:>10}" for s in systems))
    print("-" * (46 + 10 * len(systems)))
    for sec, b in by_section.items():
        label = sec if len(sec) <= 40 else sec[:39] + "\u2026"
        print(f"{label:<42}{b['count']:>4}" + "".join(f"{pct(*b[s]):>10.1f}" for s in systems))
    print("-" * (46 + 10 * len(systems)))

    overall = {s: pct(totals[s]["err"], totals[s]["n"]) for s in systems}
    print(f"{'OVERALL':<42}{len(scored):>4}" + "".join(f"{overall[s]:>10.1f}" for s in systems))

    fields = ["id", "section", "sentence", "reference_ipa"]
    for s in systems:
        fields += [f"{s}_ipa", f"{s}_per"]

    with open(output_path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows(results)
        summary_row = {"id": "OVERALL"}
        for s in systems:
            summary_row[f"{s}_per"] = overall[s]
        writer.writerow(summary_row)

    print(f"\nResults -> {output_path}")
    print("\nSummary: " + "  ".join(f"{labels[s]} {overall[s]:.1f}%" for s in systems))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="PER scoring reading all 5 columns directly, no recomputation.")
    ap.add_argument("--excel", required=True, help="Evaluation workbook (.xlsx)")
    ap.add_argument("--output", default="per_results_final.csv")
    ap.add_argument("--sheet", default=None, help="Sheet name (defaults to first sheet)")
    args = ap.parse_args()

    evaluate(args.excel, args.output, args.sheet)