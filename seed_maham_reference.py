"""
Seed Linguist Reference IPA using the proposed pipeline (EN-US fallback)
=======================================================================

Purpose
-------
This script is ONLY for preparing/initialising the linguist-reference column.

Workflow:
    1. Run the same proposed pipeline used in the evaluation:
         lexicon-first
         -> Roman numeral expansion
         -> numeric/alphanumeric handling
         -> EN-US eSpeak fallback
         -> Pakistani-English hard-rule collapses
    2. Fill the "Reference IPA (Maham)" column with that output.
    3. Give the sheet to Maham for manual verification/correction.
    4. After verification, the resulting Maham column becomes the GOLD
       reference for the actual PER evaluation.

IMPORTANT:
- The generated column is a PROVISIONAL linguist-reference seed.
- It is NOT gold until Maham verifies it.
- Lexicon entries always have priority over eSpeak.
- Add the manually transcribed pure-English words to the lexicon before
  running this script.
- This script does NOT modify any hypothesis columns.
- It only writes the reference column (column I by default).

Spreadsheet layout expected:
    A  No.
    B  ID
    C  Sentence
    D  Coverage / Key Words
    E  espeak en-us Output
    F  espeak en-pk Output
    G  Pipeline Output with en-us fallback
    H  Pipeline Output with en-pk fallback
    I  Reference IPA (Maham)

Usage:
    python seed_maham_reference.py --excel Eval.xlsx --lexicon lex.txt

Optional:
    --overwrite-reference
        By default, existing reference cells are preserved so that a
        previously verified Maham transcription is never overwritten.
        Use this flag only if you intentionally want to regenerate all
        reference seeds.

    --sheet "Evaluation Sentences"
        Change the worksheet name.

    --espeak-path "C:\\Program Files\\eSpeak NG\\espeak-ng.exe"
        Override the eSpeak executable path.
"""

import argparse
import csv
import re
import subprocess
import sys
from collections import OrderedDict

import editdistance
import openpyxl
from num2words import num2words

ESPEAK_PATH = r"C:\Program Files\eSpeak NG\espeak-ng.exe"

# Column indices (1-based) in the Evaluation Sentences sheet
COL_NO = 1
COL_ID = 2
COL_SENT = 3
COL_COVERAGE = 4
COL_ENUS = 5
COL_ENPK = 6
COL_PIPE_US = 7
COL_PIPE_PK = 8
COL_REF = 9


# --------------------------------------------------------------------------- #
# Phoneme inventory — longest-first for greedy tokenisation
# --------------------------------------------------------------------------- #

PHONEME_INVENTORY = sorted([
    # Pakistani English additional phones
    "t\u032A\u02B0", "d\u032A\u02B0", "\u0256\u02B0", "\u027D\u02B0",
    "t\u032A", "d\u032A",
    "p\u02B0", "b\u02B0", "k\u02B0", "\u0261\u02B0",
    "t\u0283\u02B0", "d\u0292\u02B0",
    "m\u02B0", "n\u02B0", "l\u02B0", "r\u02B0", "j\u02B0",
    # Nasalised vowels
    "\u0251\u0303\u02D0", "e\u0303\u02D0", "i\u0303\u02D0",
    "o\u0303\u02D0", "u\u0303\u02D0", "\u0254\u0303\u02D0", "\u00E6\u0303\u02D0",
    "\u026A\u0303", "\u028A\u0303", "\u0259\u0303",
    # Long monophthongs (including collapse targets)
    "\u0251\u02D0", "i\u02D0", "u\u02D0", "\u025C\u02D0",
    "e\u02D0", "o\u02D0", "\u0254\u02D0",
    # Remaining diphthongs
    "a\u026A", "\u0251\u026A", "a\u028A", "\u0251\u028A", "\u0254\u026A",
    # Affricates / postalveolars
    "t\u0283", "d\u0292", "\u0283", "\u0292", "\u014B",
    # Pakistani English consonants
    "\u027D", "\u0256", "q", "\u0294", "x", "\u0263",
    # Single-character phones (theta, eth, w excluded — collapsed)
    "p", "b", "t", "d", "k", "\u0261", "m", "n",
    "f", "v", "s", "z", "h", "\u0279", "j", "l", "\u027E",
    # Vowels
    "\u026A", "\u028A", "\u0259", "\u028C", "\u025B", "\u00E6", "\u0254", "\u0252",
    "\u0250", "\u0268", "\u1D7B", "\u025A", "\u025D",
    # Suprasegmentals & punctuation
    "\u02C8", "\u02CC", "\u02D0",
    ";", ":", ",", ".", "!", "?", "\u2014", "\u2026",
], key=len, reverse=True)


# --------------------------------------------------------------------------- #
# IPA tokenisation & PER
# --------------------------------------------------------------------------- #

def tokenize_ipa(ipa_string, inventory=PHONEME_INVENTORY):
    """Greedy longest-match IPA tokenisation. Unknown chars are kept as-is."""
    tokens, i = [], 0
    s = (ipa_string or "").replace(" ", "")
    while i < len(s):
        for phone in inventory:
            if s.startswith(phone, i):
                tokens.append(phone)
                i += len(phone)
                break
        else:
            tokens.append(s[i])
            i += 1
    return tokens


def phoneme_error_rate(reference_ipa, hypothesis_ipa):
    """Return (PER, n_errors, n_reference_phones)."""
    ref = tokenize_ipa(reference_ipa)
    hyp = tokenize_ipa(hypothesis_ipa)

    if not ref:
        return None, 0, 0

    dist = editdistance.eval(ref, hyp)
    return dist / len(ref), dist, len(ref)


# --------------------------------------------------------------------------- #
# eSpeak
# --------------------------------------------------------------------------- #

def g2p(text, voice="en-us"):
    """Phonemise text with eSpeak NG. Returns empty string on failure."""
    try:
        result = subprocess.run(
            [ESPEAK_PATH, "-q", "--ipa=3", "-v", voice],
            input=text,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            encoding="utf-8",
        )
        if result.returncode != 0:
            print(
                f"  [WARN] espeak {voice} failed on {text!r}: "
                f"{result.stderr.strip()}",
                file=sys.stderr,
            )
            return ""
        return " ".join(result.stdout.split())
    except Exception as exc:
        print(
            f"  [WARN] espeak {voice} failed on {text!r}: {exc}",
            file=sys.stderr,
        )
        return ""


# --------------------------------------------------------------------------- #
# Finalised Pakistani-English phonemic collapses
# --------------------------------------------------------------------------- #

def apply_pake_collapses(ipa):
    """Apply collapses ONLY to eSpeak output, never to lexicon entries."""
    ipa = ipa.replace("\u200d", "")
    ipa = ipa.replace("e\u026A", "e\u02D0").replace("\u0259\u026A", "e\u02D0")
    ipa = ipa.replace("o\u028A", "o\u02D0").replace("\u0259\u028A", "o\u02D0")
    ipa = ipa.replace("\u03B8", "t\u032A\u02B0")
    ipa = ipa.replace("\u00F0", "d\u032A")
    ipa = ipa.replace("w", "v")
    ipa = re.sub(r"t(?![\u032A\u0283\u0361])", "\u0288", ipa)   # t -> ʈ
    ipa = re.sub(r"d(?![\u032A\u0292\u0361])", "\u0256", ipa)   # d -> ɖ
    return ipa


# --------------------------------------------------------------------------- #
# Routing
# --------------------------------------------------------------------------- #

ABBREVIATIONS = {
    "mrs": "misess", "mr": "mister", "dr": "doctor", "co": "company",
    "jr": "junior", "maj": "major", "gen": "general", "drs": "doctors",
    "rev": "reverend", "lt": "lieutenant", "hon": "honorable",
    "sgt": "sergeant", "capt": "captain", "esq": "esquire",
    "ltd": "limited", "col": "colonel", "ft": "fort", "rs": "rupees",
    "st": "street",
}

ROMAN_NUMERALS = {
    "i": "one", "ii": "two", "iii": "three", "iv": "four", "v": "five",
    "vi": "six", "vii": "seven", "viii": "eight", "ix": "nine", "x": "ten",
}

PUNCTUATION = set(";:,.!?\u00A1\u00BF\u2014\u2026\"\u00AB\u00BB\u201C\u201D")

_TOKEN_RE = re.compile(r"\([ivx]+\)|\w+|[^\w\s]", re.IGNORECASE)
_PAREN_ROMAN_RE = re.compile(
    r"^\((i|ii|iii|iv|v|vi|vii|viii|ix|x)\)$", re.IGNORECASE
)
_NUM_ALPHA_RE = re.compile(r"^(\d+)([a-zA-Z]+)$")
_INTEGER_RE = re.compile(r"^\d+$")
_WORD_RE = re.compile(r"\w+")
_ABBREV_RE = {
    k: re.compile(r"\b%s\." % re.escape(k), re.IGNORECASE)
    for k in ABBREVIATIONS
}


def expand_abbreviations(text):
    for key, pattern in _ABBREV_RE.items():
        text = pattern.sub(ABBREVIATIONS[key], text)
    return text


def analyze_token(token, lexicon, fallback_voice="en-us"):
    """
    Route one token:

        lexicon -> roman -> alphanumeric -> integer -> eSpeak fallback

    Every eSpeak fallback uses the requested pipeline voice. This means
    Pipeline + EN-US and Pipeline + EN-PK differ only in their fallback voice;
    lexicon entries remain identical between the two variants.
    """
    lower = token.lower()

    if lower in lexicon:
        return lexicon[lower]

    roman = _PAREN_ROMAN_RE.match(token)
    if roman:
        return apply_pake_collapses(
            g2p(ROMAN_NUMERALS[roman.group(1).lower()], fallback_voice)
        )

    num_alpha = _NUM_ALPHA_RE.match(lower)
    if num_alpha:
        num_part, alpha_part = num_alpha.groups()

        num_ipa = apply_pake_collapses(
            g2p(num2words(int(num_part)), fallback_voice)
        )

        alpha_ipa = (
            lexicon[alpha_part]
            if alpha_part in lexicon
            else apply_pake_collapses(g2p(alpha_part, fallback_voice))
        )

        return (num_ipa + " " + alpha_ipa).strip()

    if _INTEGER_RE.match(lower):
        return apply_pake_collapses(
            g2p(num2words(int(lower)), fallback_voice)
        )

    if _WORD_RE.match(lower):
        return apply_pake_collapses(g2p(lower, fallback_voice))

    if token in PUNCTUATION:
        return token

    return ""


def pipeline_sentence(sentence, lexicon, fallback_voice="en-us"):
    """
    Run the proposed pipeline.

    The pipeline is lexicon-first. OOV words, numeric expansions, and Roman
    numeral expansions use the selected eSpeak fallback voice.
    """
    normalised = expand_abbreviations(sentence)
    tokens = _TOKEN_RE.findall(normalised)

    parts = [
        analyze_token(tok, lexicon, fallback_voice)
        for tok in tokens
    ]

    return " ".join(p for p in parts if p).strip()



# --------------------------------------------------------------------------- #
# Reference-sheet generation
# --------------------------------------------------------------------------- #

import argparse
import sys
import openpyxl


COL_NO = 1
COL_ID = 2
COL_SENT = 3
COL_REF = 9


def load_lexicon(path):
    """Load a tab-separated word<TAB>IPA lexicon, case-insensitively."""
    lexicon = {}
    with open(path, "r", encoding="utf-8") as fh:
        for line_no, line in enumerate(fh, start=1):
            line = line.rstrip("\n")
            if not line.strip() or "\t" not in line:
                continue

            word, ipa = line.split("\t", 1)
            word = word.strip().lower()
            ipa = ipa.strip()

            if not word:
                continue

            if word in lexicon and lexicon[word] != ipa:
                print(
                    f"[WARN] Duplicate lexicon entry for {word!r} "
                    f"at line {line_no}; replacing previous value.",
                    file=sys.stderr,
                )

            lexicon[word] = ipa

    return lexicon


def get_sentence_rows(ws):
    """Return spreadsheet rows that contain a sentence ID and sentence."""
    rows = []

    for row_idx, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        no_value = row[COL_NO - 1] if len(row) >= COL_NO else None
        sent_id = row[COL_ID - 1] if len(row) >= COL_ID else None
        sentence = row[COL_SENT - 1] if len(row) >= COL_SENT else None

        # Ignore section/header/blank rows.
        if no_value is not None:
            no_text = str(no_value).strip().lower()
            if no_text.startswith("section"):
                continue

        if not sent_id or not sentence:
            continue

        rows.append((row_idx, str(sent_id).strip(), str(sentence).strip()))

    return rows


def seed_reference(
    excel_path,
    lexicon_path,
    sheet_name="Evaluation Sentences",
    overwrite_reference=False,
):
    lexicon = load_lexicon(lexicon_path)

    wb = openpyxl.load_workbook(excel_path)

    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet {sheet_name!r} not found. "
            f"Available sheets: {wb.sheetnames}"
        )

    ws = wb[sheet_name]

    ws.cell(row=1, column=COL_REF, value="Reference IPA (Maham)")

    rows = get_sentence_rows(ws)

    written = 0
    preserved = 0
    failed = 0

    print(f"Loaded {len(lexicon)} lexicon entries.")
    print(f"Found {len(rows)} sentence rows.")
    print()

    for row_idx, sent_id, sentence in rows:
        existing = ws.cell(row=row_idx, column=COL_REF).value
        existing = str(existing).strip() if existing else ""

        if existing and not overwrite_reference:
            preserved += 1
            print(f"  {sent_id}  preserved existing Maham/reference value")
            continue

        # EXACTLY the proposed pipeline with EN-US as fallback.
        reference_seed = pipeline_sentence(
            sentence,
            lexicon,
            fallback_voice="en-us",
        )

        if not reference_seed:
            failed += 1
            print(
                f"  [WARN] {sent_id}  pipeline returned empty reference",
                file=sys.stderr,
            )
            continue

        ws.cell(
            row=row_idx,
            column=COL_REF,
            value=reference_seed,
        )
        written += 1
        print(f"  {sent_id}  seeded")

    wb.save(excel_path)

    print()
    print("Reference seeding complete.")
    print(f"  Newly written: {written}")
    print(f"  Preserved:     {preserved}")
    print(f"  Failed/empty:  {failed}")
    print(f"  Excel file:    {excel_path}")
    print()
    print(
        "IMPORTANT: The Reference IPA (Maham) column is only a provisional "
        "seed until Maham manually verifies/corrects it."
    )


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Seed the Maham reference column using the proposed "
            "pipeline with EN-US fallback."
        )
    )
    parser.add_argument(
        "--excel",
        required=True,
        help="Path to the evaluation Excel workbook.",
    )
    parser.add_argument(
        "--lexicon",
        required=True,
        help="Path to tab-separated word<TAB>IPA lexicon.",
    )
    parser.add_argument(
        "--sheet",
        default="Evaluation Sentences",
        help="Worksheet name (default: Evaluation Sentences).",
    )
    parser.add_argument(
        "--overwrite-reference",
        action="store_true",
        help=(
            "Overwrite existing values in column I. "
            "By default existing values are preserved."
        ),
    )
    parser.add_argument(
        "--espeak-path",
        default=None,
        help="Optional path to espeak-ng.exe.",
    )

    args = parser.parse_args()

    global ESPEAK_PATH
    if args.espeak_path:
        ESPEAK_PATH = args.espeak_path

    seed_reference(
        excel_path=args.excel,
        lexicon_path=args.lexicon,
        sheet_name=args.sheet,
        overwrite_reference=args.overwrite_reference,
    )


if __name__ == "__main__":
    main()
